from __future__ import annotations

import io
from typing import Any

import openpyxl
import polars as pl
import structlog
from openpyxl.worksheet.worksheet import Worksheet

logger = structlog.get_logger(__name__)


def read_excel_safe(raw_bytes: bytes, sheet_name: str | None = None) -> pl.DataFrame:
    """Ultra-resilient Excel parser attempting Polars fastexcel, calamine, openpyxl,

    and direct openpyxl row-extraction fallbacks.
    """
    last_err: Exception | None = None

    # Strategy 1: Polars engines (fastexcel, calamine, openpyxl)
    engines: list[str | None] = [None, "calamine", "openpyxl"]
    for engine in engines:
        try:
            kwargs: dict[str, Any] = {}
            if sheet_name is not None:
                kwargs["sheet_name"] = sheet_name
            if engine is not None:
                kwargs["engine"] = engine
            df_res: pl.DataFrame = pl.read_excel(io.BytesIO(raw_bytes), **kwargs)
            return df_res
        except Exception as e:
            last_err = e
            continue

    # Strategy 2: Direct openpyxl row-level parsing
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True)
        ws: Any = None
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active

        if ws is None or not isinstance(ws, Worksheet):
            return pl.DataFrame()

        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return pl.DataFrame()

        # Extract header and rows
        raw_header = rows[0]
        headers: list[str] = []
        for i, h in enumerate(raw_header):
            if h is None or str(h).strip() == "":
                headers.append(f"column_{i + 1}")
            else:
                headers.append(str(h).strip())

        data_rows = rows[1:]
        if not data_rows:
            # Empty table with columns
            return pl.DataFrame(schema=headers)

        return pl.DataFrame(data_rows, schema=headers, orient="row")
    except Exception as direct_err:
        logger.error("direct_openpyxl_fallback_failed", error=str(direct_err))
        if last_err is not None:
            raise ValueError(f"Failed to parse Excel file: {last_err}") from direct_err
        raise ValueError(f"Failed to parse Excel file: {direct_err}") from direct_err
